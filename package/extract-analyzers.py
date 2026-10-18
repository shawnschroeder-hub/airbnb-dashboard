# Copies the three ACTUALS workbooks into package/analyzers and dumps their main sheets to analyzers.json
import openpyxl, glob, json, shutil, os
SRC = r"C:\Users\shawn\Desktop\Agent OS\01 Projects\Airbnb Acquisitions\Analyzers"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "analyzers")
items = [("186", "423 S Rivard St", "rivard-analyzer.xlsx"), ("187", "239 W Country Ct", "country-ct-analyzer.xlsx"), ("188", "287 E Munroe St", "munroe-analyzer.xlsx")]
def fmt(c):
    v = c.value
    if isinstance(v, (int, float)):
        nf = c.number_format or ""
        if "%" in nf: return f"{v*100:.1f}%"
        if v == 0 and ("$" in nf): return "$0"
        if "$" in nf or abs(v) >= 1000: return ("-" if v < 0 else "") + f"${abs(v):,.0f}" if abs(v) >= 100 else ("-" if v < 0 else "") + f"${abs(v):,.2f}"
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    return str(v)
data = []
for num, name, short in items:
    f = glob.glob(os.path.join(SRC, num + " - *.xlsx"))[0]
    shutil.copyfile(f, os.path.join(OUT, short))
    wb = openpyxl.load_workbook(f, data_only=True)
    sheets = []
    for sh in ["Property Details", "Income and Expenses", "Deal Calculations"]:
        ws = wb[sh]; rows = []
        for r in ws.iter_rows(max_col=4):
            cells = [x for x in r if x.value not in (None, "")]
            if not cells: continue
            rows.append([fmt(x) for x in r])
        sheets.append({"name": sh, "rows": rows})
    notes = [wb["Directions and Notes"].cell(30 + i, 1).value for i in range(4)]
    data.append({"name": name, "file": "analyzers/" + short, "source": os.path.basename(f), "sheets": sheets, "notes": notes})
json.dump(data, open(os.path.join(os.path.dirname(OUT), "analyzers.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
for d in data:
    print("==", d["name"], d["file"])
    for s in d["sheets"]:
        print(" --", s["name"]); [print("   ", " | ".join(x for x in r if x not in ("None",))[:110]) for r in s["rows"]]
